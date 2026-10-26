"""Build editable Vietnamese PPTX and a PDF preview from verified project artifacts.

Charts and tables are read from CSV/JSON outputs. Narrative notes describe the
audited experiment snapshot and should be reviewed if that experiment changes.
PPTX and PDF share a layout; PDF is not rendered by Microsoft PowerPoint.
"""
from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path

import pymupdf as fitz
import pandas as pd
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

HERE = Path(__file__).resolve().parent
SRC = HERE.parent / "src"
OUT = SRC / "outputs"
W, H = 960, 540
BG, WHITE, NAVY = "F5F7FA", "FFFFFF", "132B40"
INK, MUTED, TEAL = "20384B", "607487", "008788"
BLUE, ORANGE, RED = "477AC1", "DB733E", "BD554C"
PALE, RULE = "E6F3F1", "DEE5ED"


def load(name):
    return pd.read_csv(OUT / name)


def obj(name):
    return json.loads((OUT / name).read_text())


def vn(value, decimals=0):
    return f"{float(value):,.{decimals}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def rgb(color):
    return RGBColor.from_string(color)


def emu(value):
    return Inches(value / 72)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = emu(W), emu(H)
        self.prs.core_properties.title = "Phân tích và cải tiến điểm nghẽn BPMN — Công ty X"
        self.prs.core_properties.subject = "Đồ án: BPMN, Process Mining, dự báo WIP và kịch bản cải tiến"
        self.prs.core_properties.author = "Nhóm thực hiện — bổ sung thành viên sau"
        self.prs.core_properties.language = "vi-VN"
        self.pdf = canvas.Canvas(str(HERE / "BPMN_CongTyX_30slides.pdf"), pagesize=(W, H))
        self.pdf.setTitle(self.prs.core_properties.title)
        candidates = [
            ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf", "Arial"),
            ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", "Arial"),
            ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "DejaVu Sans"),
        ]
        selected = next(((Path(a), Path(b), name) for a, b, name in candidates if Path(a).exists() and Path(b).exists()), None)
        if selected is None:
            raise FileNotFoundError("Install Arial or DejaVu Sans for Vietnamese text rendering")
        regular, bold, self.font_name = selected
        pdfmetrics.registerFont(TTFont("Body", str(regular)))
        pdfmetrics.registerFont(TTFont("Strong", str(bold)))
        self.slides, self.boxes = [], []

    def new(self, section, title, subtitle="", sources=(), dark=False):
        if self.slides:
            self.pdf.showPage()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.dark = dark
        self.record = {"slide": len(self.slides)+1, "section": section, "title": title,
                       "sources": list(sources), "visible_text": [], "notes": ""}
        self.slides.append(self.record)
        self.rect(0, 0, W, H, NAVY if dark else BG)
        self.rect(0, 0, 8, H, TEAL)
        if not dark:
            self.text(42, 24, 850, 16, section.upper(), 10, color=TEAL, bold=True)
            self.text(42, 49, 876, 44, title, 27, bold=True, color=NAVY)
            if subtitle:
                self.text(42, 96, 876, 30, subtitle, 12, color=MUTED)
        self.line(42, 502, 918, 502, "466073" if dark else RULE, .8)
        self.text(42, 510, 580, 14, "BPMN • CÔNG TY X  /  ĐỒ ÁN CUỐI KỲ", 8,
                  color="B7C6D1" if dark else MUTED)
        self.text(870, 509, 48, 15, f"{len(self.slides):02d} / 30", 9,
                  color="B7C6D1" if dark else MUTED, align="right")
        if sources:
            labels = "; ".join(s.replace("process/", "").replace("forecast/", "").replace("scenarios/", "")
                               for s in sources)
            self.text(42, 484, 875, 14, "Nguồn: " + labels, 7.5, color="B7C6D1" if dark else MUTED)
        return self

    def rect(self, x, y, w, h, fill=WHITE, stroke=None, radius=0, kind="rect"):
        styles = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE,
                  "circle": MSO_SHAPE.OVAL, "diamond": MSO_SHAPE.DIAMOND}
        shape = self.slide.shapes.add_shape(styles["round" if radius else kind], emu(x), emu(y), emu(w), emu(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if stroke:
            shape.line.color.rgb = rgb(stroke)
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        self.pdf.setFillColorRGB(*[int(fill[i:i+2], 16)/255 for i in (0, 2, 4)])
        if stroke:
            self.pdf.setStrokeColorRGB(*[int(stroke[i:i+2], 16)/255 for i in (0, 2, 4)])
        if kind == "circle":
            self.pdf.ellipse(x, H-y-h, x+w, H-y, fill=1, stroke=int(bool(stroke)))
        elif kind == "diamond":
            p = self.pdf.beginPath()
            p.moveTo(x+w/2, H-y)
            p.lineTo(x+w, H-y-h/2)
            p.lineTo(x+w/2, H-y-h)
            p.lineTo(x, H-y-h/2)
            p.close()
            self.pdf.drawPath(p, fill=1, stroke=int(bool(stroke)))
        elif radius:
            self.pdf.roundRect(x, H-y-h, w, h, radius, fill=1, stroke=int(bool(stroke)))
        else:
            self.pdf.rect(x, H-y-h, w, h, fill=1, stroke=int(bool(stroke)))
        return shape

    def line(self, x1, y1, x2, y2, color=RULE, width=1.2, arrow=False):
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2))
        shape.line.color.rgb = rgb(color)
        shape.line.width = Pt(width)
        self.pdf.setStrokeColorRGB(*[int(color[i:i+2], 16)/255 for i in (0, 2, 4)])
        self.pdf.setLineWidth(width)
        self.pdf.line(x1, H-y1, x2, H-y2)
        if arrow:
            # Arrowhead is a pair of editable connectors, shared with the PDF path.
            angle = math.atan2(y2-y1, x2-x1)
            for delta in [-.45, .45]:
                self.line(x2, y2, x2-7*math.cos(angle+delta), y2-7*math.sin(angle+delta), color, width)

    def text(self, x, y, w, h, value, size=18, bold=False, color=INK, align="left", leading=1.2):
        value = str(value)
        font = "Strong" if bold else "Body"

        def wrapped(point):
            lines = []
            for paragraph in value.split("\n"):
                words, current = paragraph.split(), ""
                if not words:
                    lines.append("")
                for word in words:
                    candidate = (current + " " + word).strip()
                    if pdfmetrics.stringWidth(candidate, font, point) <= w-3 or not current:
                        current = candidate
                    else:
                        lines.append(current)
                        current = word
                if current:
                    lines.append(current)
            return lines

        lines = wrapped(size)
        original_size = size
        while len(lines)*size*leading > h and size > 10:
            size -= .5
            lines = wrapped(size)
        if len(lines)*size*leading > h+.5:
            raise ValueError(f"Slide {len(self.slides)} text overflow: {value[:90]}")
        if any(pdfmetrics.stringWidth(s, font, size) > w+.5 for s in lines):
            raise ValueError(f"Slide {len(self.slides)} unbreakable text too wide: {value[:90]}")
        box = self.slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = False
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        for index, line in enumerate(lines):
            para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
            para.text = line
            para.font.name, para.font.size = self.font_name, Pt(size)
            para.font.bold = bold
            para.font.color.rgb = rgb(color)
            para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            para.line_spacing = Pt(size*leading)
            para.space_before = para.space_after = Pt(0)
        self.pdf.setFont(font, size)
        self.pdf.setFillColorRGB(*[int(color[i:i+2], 16)/255 for i in (0, 2, 4)])
        for index, line in enumerate(lines):
            tx = x
            if align == "right":
                tx = x+w-pdfmetrics.stringWidth(line, font, size)
            elif align == "center":
                tx = x+(w-pdfmetrics.stringWidth(line, font, size))/2
            self.pdf.drawString(tx, H-y-size*.91-index*size*leading, line)
        self.record["visible_text"].append(value)
        self.boxes.append({"slide": len(self.slides), "x": x, "y": y, "w": w, "h": h,
                           "font_size": size, "requested_size": original_size, "text": value})
        return box

    def card(self, x, y, w, h, title, body, accent=TEAL, size=17):
        self.rect(x, y, w, h, WHITE, radius=8)
        self.rect(x, y, 4, h, accent)
        if h <= 140:
            self.text(x+16, y+10, w-30, 28, title, 18, bold=True, color=accent)
            self.text(x+16, y+43, w-32, h-49, body, size)
        else:
            self.text(x+16, y+15, w-30, 38, title, 19, bold=True, color=accent)
            self.text(x+16, y+61, w-32, h-73, body, size)

    def metric(self, x, y, w, value, label, hint="", accent=TEAL):
        self.rect(x, y, w, 129, WHITE, radius=8)
        self.text(x+15, y+12, w-28, 48, value, 36, bold=True, color=accent)
        self.text(x+15, y+67, w-28, 28, label, 16, bold=True)
        if hint:
            self.text(x+15, y+101, w-28, 21, hint, 10.5, color=MUTED)

    def band(self, value, y=438, fill=PALE, color=TEAL):
        self.rect(42, y, 876, 35, fill, radius=5)
        self.text(55, y+8, 850, 21, value, 13, bold=True, color=color)

    def table(self, x, y, widths, headers, rows, row_h=43, size=14, highlights=()):
        total = sum(widths)
        self.rect(x, y, total, row_h, NAVY)
        cursor = x
        for width, header in zip(widths, headers):
            self.text(cursor+10, y+8, width-18, row_h-12, header, size, bold=True, color=WHITE)
            cursor += width
        for i, row in enumerate(rows):
            yy = y+(i+1)*row_h
            self.rect(x, yy, total, row_h, PALE if i in highlights else WHITE if i % 2 == 0 else "EDF1F6")
            cursor = x
            for width, cell in zip(widths, row):
                self.text(cursor+10, yy+9, width-18, row_h-13, str(cell), size, bold=i in highlights)
                cursor += width

    def bars(self, x, y, w, labels, values, max_value=None, colors=None, value_format=None, row_h=43, label_w=160):
        maximum = max_value or max(values)*1.12 or 1
        for i, (label, value) in enumerate(zip(labels, values)):
            yy = y+i*row_h
            self.text(x, yy+1, label_w-12, row_h-6, label, 14)
            start, span = x+label_w, w-label_w-72
            self.rect(start, yy+3, span, 20, "E5EAF0", radius=3)
            self.rect(start, yy+3, max(1, span*value/maximum), 20,
                      colors[i] if colors else TEAL, radius=3)
            self.text(start+span+10, yy, 65, 27, value_format(value) if value_format else vn(value, 0),
                      14, bold=True, color=colors[i] if colors else TEAL)

    def notes(self, value):
        self.record["notes"] = value
        self.slide.notes_slide.notes_text_frame.text = (
            f"SLIDE {len(self.slides):02d} — {self.record['title']}\n\n{value}\n\n"
            "Nguồn kết quả (tương đối với src/outputs):\n" + "\n".join(self.record["sources"])
        )

    def save(self):
        assert len(self.slides) == 30
        self.prs.save(HERE / "BPMN_CongTyX_30slides.pptx")
        self.pdf.save()
        source_paths = sorted({source for record in self.slides for source in record["sources"]})
        (HERE / "slide_sources.json").write_text(json.dumps({
            "slides": self.slides,
            "sources": [{"artifact": p, "sha256": sha(OUT / p)} for p in source_paths],
            "preview": "PDF shares layout primitives with PPTX; not a Microsoft Office renderer",
        }, ensure_ascii=False, indent=2), encoding="utf-8")
        (HERE / "layout_audit.json").write_text(json.dumps({"text_boxes": self.boxes}, ensure_ascii=False,
                                                          indent=2), encoding="utf-8")
        notes = "# Ghi chú thuyết trình — 30 slide\n\nTên thành viên để bổ sung sau.\n\n"
        for slide in self.slides:
            notes += f"## {slide['slide']:02d}. {slide['title']}\n\n{slide['notes']}\n\n"
            if slide["sources"]:
                notes += "Nguồn: " + ", ".join(f"`src/outputs/{p}`" for p in slide["sources"]) + "\n\n"
        (HERE / "speaker_notes.md").write_text(notes, encoding="utf-8")


def build():
    status, verification = obj("metadata/run_all.json"), obj("metadata/verification.json")
    if status["status"] != "success" or status["quick"] or verification["status"] != "passed":
        raise ValueError("A verified full pipeline run is required")
    cases = load("process/case_metrics.csv")
    evidence = load("process/bottleneck_evidence.csv")
    timing = load("process/task_timing.csv")
    qc = load("process/qc_by_layer.csv")
    repeats = load("process/paint_repeated_visits.csv")
    pareto = load("process/defect_pareto.csv")
    model = load("forecast/model_comparison.csv")
    ensemble = load("forecast/ensemble_metrics_month_bootstrap.csv")
    gap = load("scenarios/gap_replay_summary.csv")
    workload = load("scenarios/repeat_workload_sensitivity.csv")
    timers = load("scenarios/qc7_timer_dispatch_replay.csv")
    audit = obj("audit/summary.json")
    process = obj("process/summary.json")
    contract = obj("forecast/feature_contract.json")
    d = Deck()

    # 01 — Cover, editable member placeholder.
    d.new("Đồ án cuối kỳ", "Điểm nghẽn BPMN", dark=True)
    d.text(48, 48, 820, 22, "HOẠCH ĐỊNH NGUỒN LỰC DOANH NGHIỆP NÂNG CAO", 12, color="A1C9D0", bold=True)
    d.text(48, 123, 800, 134, "TÌM ĐIỂM NGHẼN\nTRONG QUY TRÌNH BPMN", 39, bold=True, color=WHITE)
    d.text(50, 283, 785, 63, "Phân tích quy trình sản xuất Công ty X\nvà đánh giá các kịch bản cải tiến", 24, color="C2D9E1")
    d.rect(50, 373, 525, 40, TEAL, radius=5)
    d.text(65, 383, 495, 22, "BPMN  →  PROCESS MINING  →  ML  →  CẢI TIẾN", 14, bold=True, color=WHITE)
    d.text(50, 450, 820, 27, "Thành viên: [Bổ sung sau]", 16, color=WHITE)
    d.notes("Giới thiệu bài toán thực tế của Công ty X đã ẩn danh. Nhóm sử dụng dữ liệu đã có và code tái lập để đi từ dấu vết vận hành đến phương án cải tiến. ML là nhánh thực nghiệm hỗ trợ; kết quả cải tiến là kịch bản có giả định, chưa phải hiệu quả triển khai tại doanh nghiệp. Điền tên thành viên vào textbox trên trang bìa khi chốt nhóm.")

    # 02 — Lead with the answer.
    d.new("01 / Bài toán", "Ba kết luận định hướng toàn bộ đồ án",
          "Tách rõ điều dữ liệu cho thấy, điều mô hình dự báo và điều còn phải thử nghiệm.",
          ["process/bottleneck_evidence.csv", "forecast/model_comparison.csv"])
    d.card(42, 151, 280, 249, "Ưu tiên chà nhám–sơn", "Có đồng thời tồn WIP cao, nhiều lượt làm lại và thời gian giữa các bước kéo dài. Đây là vùng cần điều tra trước.")
    d.card(340, 151, 280, 249, "Baseline vẫn rất mạnh", "RF WIP đạt F1 ≈ 0,608 cho nhãn ùn WIP mới. Graph + Node2Vec chưa cải thiện kết quả.", BLUE)
    d.card(638, 151, 280, 249, "Cải tiến có điều kiện", "Đánh giá riêng giảm khoảng cách, giảm workload làm lại và bố trí QC7; chưa suy ra mức giảm CT toàn nhà máy.", ORANGE)
    d.band("Đầu ra của đồ án: bằng chứng có thể truy nguyên + phương án có thể kiểm chứng.")
    d.notes("Bắt đầu bằng ba kết luận để người nghe hiểu lý do của các bước phía sau. Bottleneck ở đây được dùng như điểm ưu tiên điều tra dựa trên nhiều chỉ dấu, chưa phải nút giới hạn công suất đã được đo utilization. F1 0,608 là trung bình 5 seed của RF WIP cho nhãn calendar_v1; không phải điểm trong Excel E1 cũ. Các lợi ích To-Be được đánh giá có điều kiện.")

    # 03 — Research questions.
    d.new("01 / Bài toán", "Bốn câu hỏi cần trả lời bằng code",
          "Mỗi câu hỏi có đầu vào, phép đo và sản phẩm kiểm chứng tương ứng.")
    for i, (title, body) in enumerate([
        ("01  Chậm ở đâu?", "Đo khoảng thời gian quan sát, khoảng cách giữa sự kiện, WIP và vòng lặp."),
        ("02  Vì sao đáng ưu tiên?", "Đối chiếu thời gian, lỗi theo lớp và số liệu tổng thể; tách bằng chứng khỏi giả thuyết."),
        ("03  Có cảnh báo trước được không?", "Dự báo ùn WIP tháng sau, so với quy tắc đơn giản và RF trước khi thêm graph."),
        ("04  Thay đổi nào có tiềm năng?", "Định nghĩa To-Be cụ thể, tính chi phí can thiệp và phân tích độ nhạy giả định."),
    ]):
        d.card(42+(i % 2)*447, 145+(i//2)*140, 429, 124, title, body, [TEAL, BLUE, ORANGE, TEAL][i], 15)
    d.notes("Phân biệt ba tầng công việc: chẩn đoán mô tả, dự báo và đánh giá kịch bản. Câu hỏi nguyên nhân phải đi kèm xác minh nghiệp vụ, không dùng tương quan hay vector để kết luận nhân quả. Pipeline không xây ERP hoặc frontend; nguồn lực được tập trung vào dữ liệu, kiểm chứng và kết quả có thể dùng cho báo cáo.")

    # 04 — Learning-to-application map.
    d.new("01 / Nền tảng", "Kiến thức đã học được chuyển thành nhiệm vụ cụ thể",
          "Áp dụng đúng câu hỏi; không bắt buộc dùng mọi mô hình đã học.")
    d.table(42, 143, [225, 360, 291], ["Kiến thức", "Ứng dụng trong đồ án", "Sản phẩm code"], [
        ["BPMN & hệ thống công việc", "Phạm vi, vai trò, QC, nhánh hủy / sửa", "BPMN + ánh xạ MES"],
        ["Cycle Time & rework", "Thời gian quan sát, lượt lặp theo lớp", "Case / task metrics"],
        ["Event log & mạng bàn giao", "Chuẩn hóa, lỗi thời gian, chuyển người", "Audit + bảng truy nguyên"],
        ["Process Mining / Petri net", "DFG, discovery, replay mô hình", "PNML / XES / diagnostics"],
        ["Embedding & phân loại", "Node2Vec + RF, đối chiếu baseline", "Holdout theo thời gian"],
        ["BPR / tái thiết kế", "Cơ chế cải tiến và độ nhạy giả định", "To-Be + scenario results"],
    ], row_h=43, size=13)
    d.notes("Bản tổng hợp 6 buổi cung cấp chuỗi kiến thức từ mô hình hóa đến thời gian, nhật ký, tái thiết kế và biểu diễn graph. Các bài Buổi 7 cung cấp thêm thực hành embedding/classification. Đồ án chọn Node2Vec để có phần thực nghiệm graph phù hợp dữ liệu trạm–tháng. Không ép Word2Vec, Doc2Vec, Graph2Vec hoặc GNN vào khi dữ liệu và câu hỏi chưa cần. Công thức rework được dùng với phạm vi tương đương và giả định minh bạch.")

    # 05 — Production overview.
    d.new("02 / Hiểu quy trình", "Một serial đi qua nhiều công đoạn và nhiều cổng QC",
          "Bản đồ tổng quan theo mô tả nghiệp vụ; hình dưới là sơ đồ diễn giải, không thay thế BPMN chi tiết.")
    labels = ["Chuẩn bị &\ntạo hình", "Gia công &\ndán keo", "Chà nhám &\nsơn theo lớp", "Hoàn thiện &\nkiểm tra cuối", "Đóng gói &\nnhập kho"]
    for i, label in enumerate(labels):
        x = 42+i*178
        d.rect(x, 175, 164, 102, TEAL if i == 2 else WHITE, stroke=None, radius=7)
        d.text(x+12, 194, 140, 64, label, 19, bold=True, color=WHITE if i == 2 else INK, align="center")
        if i < 4:
            d.line(x+165, 225, x+176, 225, MUTED, 1.4, arrow=True)
    d.card(42, 314, 429, 104, "Các kiểm tra bắt buộc", "QC0–QC8 và kiểm tra cuối theo tuyến; có sửa được, đạt hoặc hủy.", BLUE, 15)
    d.card(489, 314, 429, 104, "Vùng phân tích trọng tâm", "Chà nhám–sơn–QC; giữ liên hệ với WIP và các bước trước/sau.", TEAL, 15)
    d.notes("Tóm tắt tài liệu Mô tả quy trình.docx và các BPMN Hình 1–5. Việc có CNC và có QC2 là hai quyết định độc lập. Tạo hình có yêu cầu may và loại bỏ ba via tùy tuyến. Hoàn thiện có QC8, đánh bóng, chờ ổn định và QC7 tùy sản phẩm. File Tong the BPMN kết thúc ở sản phẩm đạt yêu cầu, trong khi Hình 1 còn có đóng gói/nhập kho; chênh lệch phạm vi này được giữ làm diagnostic, không âm thầm sửa log.")

    # 06 — Explain loops without assuming rework.
    d.new("02 / Hiểu quy trình", "Hai kiểu quay vòng phải được phân biệt",
          "Cùng đi lại một vùng quy trình, nhưng ý nghĩa nghiệp vụ và cách đếm khác nhau.")
    for i, (x, label) in enumerate([(42, "Chà nhám"), (260, "Chuẩn bị bề mặt"), (478, "Sơn lớp k"), (696, "QC6")]):
        d.rect(x, 176, 184, 64, WHITE, stroke=RULE, radius=6)
        d.text(x+10, 194, 164, 32, label, 18, bold=True, align="center")
        if i < 3:
            d.line(x+185, 208, x+214, 208, MUTED, 1.4, arrow=True)
    d.line(786, 242, 786, 283, RED, 2)
    d.line(786, 283, 132, 283, RED, 2)
    d.line(132, 283, 132, 242, RED, 2, arrow=True)
    d.text(279, 260, 410, 20, "QC6 không đạt, sửa được → làm lại lớp k", 13, bold=True, color=RED, align="center")
    d.card(42, 326, 429, 103, "Làm lại trong cùng lớp", "Khóa đếm: case + task + lớp. Lượt bổ sung là ứng viên rework.", RED, 14)
    d.card(489, 326, 429, 103, "Chuyển sang lớp kế tiếp", "QC6 đạt và còn lớp: quay lại cho lớp k+1. Đây là công việc theo tuyến.", TEAL, 14)
    d.notes("Sơ đồ là phần chiếu đơn giản để giải thích vòng lặp, đã lược QC chà nhám và nhánh che chắn cho dễ nhìn; BPMN thực vẫn giữ các kiểm tra. Trong code, repeated visits được đếm theo case/task/layer. Một lượt bổ sung chưa luôn chứng minh lỗi làm lại, vì trace có thể thiếu hoặc có nghiệp vụ khác, nên thuật ngữ đúng là lượt bổ sung quan sát được trong cùng lớp. Chuyển lớp là công việc dự kiến, không gán nhãn defect.")

    # 07 — Source hierarchy.
    d.new("03 / Dữ liệu", "Ba tầng dữ liệu bổ sung cho nhau",
          "Không xem 8 serial là đại diện cho nhà máy; không xem mọi dòng trạm–tháng là độc lập.",
          ["audit/summary.json", "forecast/feature_contract.json", "data/population_rework.csv"])
    d.metric(42, 146, 280, vn(audit["events"]), "Sự kiện MES", "8 case • 7 sản phẩm • 62 activity")
    d.metric(340, 146, 280, "242.912", "Serial hoàn thành", "Tổng thể làm lại B1 • 20 tháng", BLUE)
    d.metric(638, 146, 280, "54 × 26", "Trạm × kỳ có WIP", "1.404 dòng snapshot được cung cấp", ORANGE)
    d.table(42, 303, [200, 340, 336], ["Tầng", "Vai trò", "Giới hạn chính"], [
        ["A1 / MES", "Truy vết thời gian, lỗi, lớp sơn", "Mẫu chọn có chủ đích; trace chưa đủ cuối"],
        ["B1 / B3 / C1", "Đối chiếu làm lại, hủy, stock tổng thể", "Khác kỳ và khác mẫu số"],
        ["D1 / D2 / D4", "Thực nghiệm graph và WIP tháng sau", "Thiếu raw/code tạo bảng nguồn"],
    ], row_h=38, size=12.5)
    d.notes("Số MES đọc trực tiếp từ A1: 1.397 sự kiện, 8 case, 7 sản phẩm, 62 activity, 57 mã trạm. B1 tổng hợp số vòng làm lại trên 242.912 serial thành phẩm, theo ghi chú B0. WIP có 26 kỳ quan sát giữa 01/2019 và 04/2021, thiếu hoặc loại hai tháng. D4 là snapshot dẫn xuất được cung cấp, không được tái tạo từ A1. Các tầng được đối chiếu theo công đoạn; không ghép nhầm các ID ẩn danh khác hệ hoặc mẫu số khác nhau.")

    # 08 — Audit.
    d.new("03 / Dữ liệu", "Audit trước khi chạy thuật toán",
          "Giữ dấu vết lỗi và quy tắc xử lý; không xóa dữ liệu chỉ để làm đẹp độ đo.", ["audit/summary.json", "audit/excel_control_differences.csv"])
    d.table(42, 146, [305, 110, 461], ["Phát hiện", "Số lượng", "Quyết định xử lý"], [
        ["Complete trước Start", str(audit["issues"]["NEGATIVE_DURATION"]), "Giữ sự kiện; để thiếu thời lượng phân tích"],
        ["Task chưa khớp phạm vi BPMN", str(audit["issues"]["UNMAPPED_TASK"]), "Giữ log-only / thiếu ánh xạ, xác minh nghiệp vụ"],
        ["Dòng thiếu định mức", str(audit["issues"]["MISSING_STANDARD"]), "Không tự điền số; nêu rõ mẫu số"],
        ["Timer QC7 ngắn hơn 72 giờ", str(audit["issues"]["TIMER_UNDER_72H"]), "Kiểm tra riêng timestamp BLOCK/UNBLOCK"],
        ["SCRAP rồi RECOVER cùng trạm", "1 cặp", "Bỏ cả hai trạng thái khỏi kiểm tra luồng"],
    ], row_h=48, size=14)
    d.band("Mỗi sự kiện giữ event_id + source_row để truy ngược về A1.")
    d.notes("Hai lỗi thời gian thuộc CASE05, chưa tự sửa đồng hồ AM/PM. Năm sự kiện không có node tương ứng trong mô hình đang dùng gồm hai task thực chưa có trong BPMN và ba sự kiện đóng gói ngoài phạm vi file tổng thể. Bảng A6 cũ giữ dòng SCRAP dù A5 yêu cầu bỏ cả SCRAP/RECOVER; code áp dụng quy tắc đã thống nhất và xuất bảng khác biệt. Thiếu định mức không đồng nghĩa thời lượng bằng 0. Tất cả dữ liệu nguồn được giữ nguyên.")

    # 09 — Architecture.
    d.new("04 / Phương pháp", "Một pipeline tái lập từ dữ liệu đến quyết định",
          "Mỗi bước xuất artifact kiểm tra được; kết quả Excel cũ chỉ dùng để đối chiếu.", ["metadata/verification.json"])
    steps = [("01", "Đọc & audit", "data.py"), ("02", "Đo & truy nguyên", "analysis.py"),
             ("03", "Khám phá luồng", "mining.py"), ("04", "Dự báo WIP", "features / models"),
             ("05", "Đánh giá To-Be", "simulation.py")]
    for i, (number, label, module) in enumerate(steps):
        x = 42+i*178
        d.rect(x, 157, 164, 157, WHITE, radius=7)
        d.text(x+13, 169, 130, 39, number, 28, color=TEAL, bold=True)
        d.text(x+13, 220, 138, 39, label, 18, bold=True)
        d.text(x+13, 278, 138, 22, module, 11, color=MUTED)
        if i < 4:
            d.line(x+165, 235, x+176, 235, MUTED, 1.5, arrow=True)
    d.card(42, 344, 429, 84, "Artifact", "CSV/JSON, XES/PNML, model, hình PNG/SVG và notebook demo.", BLUE, 14)
    d.card(489, 344, 429, 84, "Tái lập", "uv.lock, config, hash nguồn và kiểm tra lại model đã lưu.", TEAL, 14)
    d.notes("Tất cả code chạy bằng uv, có môi trường khóa phiên bản và config tách khỏi logic. A6–A11/E1–E4 không làm kết quả mặc định. Pipeline xuất các bảng so sánh khác biệt với Excel cũ, tạo metadata về nguồn và version nhãn. Verify tính lại 64 nhóm độ đo và tái tạo dự đoán từ các model đã lưu. Slide chỉ đọc các kết quả đã có, không tạo thêm dữ liệu doanh nghiệp.")

    # 10 — Measurements.
    d.new("04 / Phương pháp", "Đo đúng thời gian trước khi gọi tên điểm nghẽn",
          "Log quét trạm chưa có ready_time, nên chưa tách chính xác hàng chờ nguồn lực.")
    d.text(42, 147, 876, 30, "Complete trước  →  Start hiện tại  →  Complete hiện tại", 23, bold=True, color=NAVY, align="center")
    d.rect(80, 211, 380, 54, PALE, radius=6)
    d.rect(477, 211, 401, 54, "E6EDF8", radius=6)
    d.text(90, 228, 360, 29, "Khoảng cách giữa hai lần quét", 18, bold=True, color=TEAL, align="center")
    d.text(487, 228, 381, 29, "Thời gian ghi nhận tại trạm", 18, bold=True, color=BLUE, align="center")
    d.card(42, 306, 280, 112, "Khoảng quan sát case", "max(Complete) − min(Start). Chưa xác nhận tất cả case đã hoàn tất.", TEAL, 14)
    d.card(340, 306, 280, 112, "Khoảng cách quét", "Start hiện tại − Complete trước. Có thể gồm chờ, vận chuyển, lịch ca.", BLUE, 14)
    d.card(638, 306, 280, 112, "Định mức / thời gian", "Là proxy so sánh theo BOM; không coi là thời gian tạo giá trị đo trực tiếp.", ORANGE, 14)
    d.notes("Code giữ source order theo MES và tách raw duration/gap với giá trị phân tích hợp lệ. Nếu Complete−Start âm, sự kiện vẫn được giữ cho quan sát luồng còn duration được để thiếu. Khoảng thời gian ở trạm có thể chứa sấy, đổi ca hoặc chờ nên không gọi là pure service time. Khi thiếu ready_time không thể tách queue time chuẩn. Các công thức trên dùng giờ lịch, không quy đổi sang giờ làm việc vì chưa có working calendar.")

    # 11 — Case spans.
    d.new("05 / Bằng chứng", "Khoảng quan sát giữa các case chênh lệch rất lớn",
          "Sự phân tán quan trọng hơn một giá trị trung bình duy nhất.", ["process/case_metrics.csv"])
    d.bars(42, 146, 584, cases.case_id.tolist(), cases.observed_span_days.tolist(), row_h=34,
           label_w=90, colors=[ORANGE if v in ["CASE01", "CASE07"] else TEAL for v in cases.case_id],
           value_format=lambda v: vn(v, 1))
    d.text(135, 426, 460, 19, "Đơn vị: ngày theo lịch; không phải giờ gia công thuần", 10.5, color=MUTED)
    d.metric(660, 152, 258, vn(cases.observed_span_days.median(), 1), "Ngày — trung vị", "8 case được chọn có chủ đích")
    d.metric(660, 298, 258, vn(cases.observed_span_days.mean(), 1), "Ngày — trung bình", "Một số trace thiếu bước kết thúc", ORANGE)
    d.notes("Khoảng quan sát trung vị khoảng 46,4 ngày và trung bình khoảng 79,5 ngày. CASE01 khoảng 134,7 ngày, CASE07 khoảng 218,8 ngày. CASE02/CASE07 kết thúc trong vùng chuẩn bị bề mặt, CASE03 ở QC6 FAIL; do đó không coi các mốc này là thời gian hoàn thành sản phẩm. Mục tiêu biểu đồ là nhận diện dị biệt và chọn case kiểm tra, không xếp hạng hiệu quả sản phẩm hoặc nhân viên.")

    # 12 — Triangulation.
    d.new("05 / Bằng chứng", "Chà nhám–sơn nổi bật khi đối chiếu nhiều nguồn",
          "Stock, làm lại và thời gian là các chỉ dấu khác nhau; không cộng thành một điểm số tùy ý.", ["process/bottleneck_evidence.csv"])
    subset = evidence.head(5)
    rows = [[r.stage.replace("Tạo hình/Tách khuôn/Ổn định", "Tạo hình / ổn định"), vn(r.population_month_end_wip_mean),
             "—" if pd.isna(r.population_mean_repeat_operations) else "<0,001" if 0 < r.population_mean_repeat_operations < .001 else vn(r.population_mean_repeat_operations, 3),
             "Ưu tiên điều tra" if r.screening_category == "high_stock_and_rework" else "Theo dõi stock"]
            for r in subset.itertuples()]
    d.table(42, 150, [302, 180, 178, 216], ["Công đoạn", "WIP TB cuối tháng", "Lượt làm lại / SP", "Nhận định"], rows,
            row_h=48, size=13, highlights=[1, 3])
    d.band("Sàng lọc: thuộc top 5 stock và R ≥ 0,5 → vùng có đồng thời stock cao và làm lại nhiều.")
    d.notes("Chà nhám có WIP trung bình 3.357,35 SP, R=1,1986; sơn có WIP 1.988,96 SP, R=1,2076. Vận chuyển/chờ chuyển có stock cao nhất, nhưng chưa gán được cho công đoạn sau và không có R tương ứng. Quy tắc top 5 cộng R≥0,5 là sàng lọc mô tả minh bạch, không phải nhãn bottleneck được huấn luyện hay định luật công suất. Nên ưu tiên vùng chà nhám–sơn để xác minh nguyên nhân và thiết kế thử nghiệm.")

    # 13 — Task timing.
    d.new("05 / Bằng chứng", "Khoảng cách trước chà nhám là thành phần lớn trong mẫu",
          "Phân rã theo task để tránh nhầm lưu tại trạm với thời gian xử lý thuần.", ["process/task_timing.csv", "process/summary.json"])
    chosen = timing.head(5)
    short = ["Chà nhám", "Đánh bóng", "Sơn", "Ngoài phạm vi BPMN", "QC chà nhám"]
    max_total = (chosen.gap_hours+chosen.residence_hours).max()
    for i, (r, label) in enumerate(zip(chosen.itertuples(), short)):
        y = 154+i*48
        d.text(42, y+2, 170, 31, label, 16)
        start, span = 224, 535
        d.rect(start, y+4, span*r.gap_hours/max_total, 24, TEAL)
        d.rect(start+span*r.gap_hours/max_total, y+4, max(1, span*r.residence_hours/max_total), 24, ORANGE)
        d.text(775, y+1, 143, 35, vn(r.gap_hours+r.residence_hours)+" giờ", 15, bold=True)
    d.rect(226, 414, 12, 12, TEAL)
    d.text(244, 411, 285, 20, "Khoảng cách giữa các lần quét", 11)
    d.rect(540, 414, 12, 12, ORANGE)
    d.text(558, 411, 290, 20, "Thời gian ghi nhận tại trạm", 11)
    d.band(f"{vn(process['inter_scan_gap_to_span_ratio']*100,1)}% tổng khoảng quan sát là gap ghi nhận; chưa phải tỷ lệ chờ thuần.")
    d.notes("Ở mức task, chà nhám có khoảng 6.105,5 giờ gap và 106,7 giờ lưu trạm; sơn có khoảng 810,2 giờ gap và 2.737,6 giờ lưu trạm. Tổng này trên nhiều case, không phải thời gian của một sản phẩm. Trạm sơn CASE01 lưu khoảng 95 ngày ảnh hưởng lớn phần màu cam. Mẫu có phần lớn scan rất ngắn nên chênh lệch giữa quét là proxy hữu ích, nhưng không đủ để suy ra utilization hoặc queue resource. 79,3% là tỷ lệ gap theo định nghĩa của code.")

    # 14 — Outliers.
    d.new("05 / Kiểm tra độ vững", "Ngoại lai thay đổi mức độ, nên phải đọc cả độ nhạy",
          "Chà nhám vẫn đứng đầu về tổng gap trong các tập con được kiểm tra.", ["process/timing_sensitivity.csv"])
    sens = load("process/timing_sensitivity.csv")
    configs = [("all_events", "Toàn bộ 8 case"), ("exclude_case01", "Bỏ CASE01"),
               ("exclude_case07", "Bỏ CASE07"), ("exclude_both_long_cases", "Bỏ cả hai")]
    values = [sens[(sens.subset == code) & sens.activity.eq("Chà nhám lớp được chỉ định")].gap_hours.iloc[0]
              for code, _ in configs]
    d.bars(42, 163, 565, [label for _, label in configs], values, row_h=56, label_w=154,
           colors=[TEAL, BLUE, ORANGE, RED], value_format=lambda v: vn(v))
    d.card(648, 151, 270, 255, "Cách diễn giải", "CASE07 làm tổng gap chà nhám tăng mạnh.\n\nKhông xóa case dài mặc định; giữ cả bản gốc và phân tích độ nhạy.\n\nƯu tiên điều tra còn được hỗ trợ bởi WIP và R tổng thể.", ORANGE, 16)
    d.band("Không dùng kết quả tổng mẫu để suy ra thời gian chờ điển hình của toàn nhà máy.")
    d.notes("Tổng gap chà nhám lần lượt khoảng 6.105,5 giờ; 6.024,1 giờ khi bỏ CASE01; 1.406,6 giờ khi bỏ CASE07; 1.325,2 giờ khi bỏ cả hai. Đây là so sánh tổng trên tập con có số case khác nhau, không phải phép ước lượng causal impact hay trung bình chuẩn hóa. Code cũng xuất số case và event trong từng tập con để kiểm tra. Mục đích là biết kết luận nào bị một vài case chi phối và tránh overclaim.")

    # 15 — QC by layer.
    d.new("05 / Chất lượng", "Lỗi QC6 trong mẫu tập trung ở lớp 3 và lớp 4",
          "Đếm theo lượt kiểm tra; không đồng nhất với tỷ lệ sản phẩm lỗi toàn nhà máy.", ["process/qc_by_layer.csv", "process/qc_outcomes.csv"])
    q = qc[qc.bpmn_activity.str.startswith("QC6")]
    for i, r in enumerate(q.itertuples()):
        y = 158+i*55
        d.text(42, y, 110, 30, f"Lớp {int(r.paint_layer)}", 17, bold=True)
        span, start = 425, 157
        d.rect(start, y+2, span*r.passes/117, 27, TEAL)
        d.rect(start+span*r.passes/117, y+2, max(1, span*r.failures/117), 27, RED)
        d.text(597, y+1, 320, 32, f"{int(r.passes)} đạt / {int(r.failures)} lỗi  ({int(r.checks)} lượt)", 16)
    d.rect(42, 392, 876, 81, PALE, radius=7)
    d.text(60, 406, 240, 45, vn(q.passes.sum()/q.checks.sum()*100,1)+"% PASS", 27, bold=True, color=TEAL)
    d.text(326, 406, 565, 46, "176/176 lần FAIL tại QC6 nằm ở lớp 3–4. Đây là bằng chứng để khoanh vùng kiểm tra nguyên nhân.", 17)
    d.notes("QC6 có 226 lượt kiểm tra hợp lệ PASS/FAIL-DEFECT: 50 đạt và 176 không đạt. Lớp 1: 8/8 đạt; lớp 2: 10/10 đạt; lớp 3: 8 đạt và 83 lỗi; lớp 4: 24 đạt và 93 lỗi. Tất cả 176 lỗi QC6 trong mẫu nằm ở lớp 3–4. Mẫu được chọn vì nhiều làm lại, nên không dùng tỷ lệ 22,1% PASS làm tỷ lệ trung bình nhà máy hoặc mục tiêu áp dụng cho mọi sản phẩm. Các nguyên nhân phòng sơn/độ nhớt/độ dày cần được xác minh, chưa có trong log.")

    # 16 — Native heatmap.
    d.new("05 / Làm lại", "225 lượt sơn trên 30 tổ hợp case–lớp quan sát được",
          "195 lượt bổ sung trong cùng lớp; không tính lớp kế tiếp là làm lại.", ["process/paint_repeated_visits.csv", "process/summary.json"])
    pivot = repeats.pivot(index="case_id", columns="paint_layer", values="observed_visits")
    for col in range(4):
        d.text(184+col*95, 142, 88, 25, f"Lớp {col+1}", 14, bold=True, align="center")
    for i, (case, row) in enumerate(pivot.iterrows()):
        yy = 178+i*33
        d.text(42, yy+5, 130, 24, case, 15, bold=True)
        for col in range(4):
            value = row.get(col+1, float("nan"))
            color = "E7ECF2" if pd.isna(value) else "E4F1EF" if value <= 2 else "9DCCCA" if value < 15 else TEAL
            d.rect(184+col*95, yy, 86, 28, color, radius=3)
            d.text(188+col*95, yy+4, 78, 22, "—" if pd.isna(value) else str(int(value)), 15,
                   bold=True, color=WHITE if color == TEAL else INK, align="center")
    d.card(615, 151, 303, 258, "Một ví dụ rõ", "CASE02 có 43 lượt sơn ở lớp 3.\n\nCASE03 có 41 lượt ở lớp 3.\n\nĐếm lượt bổ sung giúp định lượng workload; chưa chứng minh mọi lượt đều do cùng một lỗi.", TEAL, 17)
    d.band("Ô trống: không quan sát lớp đó trong case. Số trong ô = lượt sơn quan sát, gồm lượt đầu.")
    d.notes("Nhóm không lấy 225 trừ một lượt cho mỗi case, vì sản phẩm cần nhiều lớp. Thay vào đó có 30 tổ hợp case–lớp quan sát được, lấy tổng max(visits−1,0) bằng 195. Tỷ số 225/30 bằng 7,5 là bội số lượt trên một lượt đầu cho mỗi case–lớp được quan sát; không tự gọi đây là 7,5 lần tuyến chuẩn đầy đủ khi trace có thể thiếu. Mỗi dòng dữ liệu có danh sách event_id để truy nguyên.")

    # 17 — Pareto and cause hypotheses.
    d.new("05 / Nguyên nhân cần kiểm chứng", "Ba nhóm lỗi chiếm 90% số lần ghi nhận lỗi",
          "Pareto giúp thu hẹp phạm vi điều tra; không tự chỉ ra nguyên nhân gốc.", ["process/defect_pareto.csv"])
    labels = ["Vết lõm / co ngót", "Bề mặt không đều", "Lỗ kim"]
    d.bars(42, 164, 560, labels, pareto.head(3)['count'].tolist(), row_h=64, label_w=195,
           colors=[TEAL, BLUE, ORANGE])
    d.text(42, 372, 560, 49, f"{vn(pareto.head(3)['count'].sum())} / {vn(pareto['count'].sum())} lần lỗi  =  90%", 26,
           bold=True, color=TEAL)
    d.card(640, 148, 278, 274, "Cần kiểm chứng thêm", "Điều kiện bề mặt đầu vào?\n\nĐộ nhớt, độ dày lớp sơn?\n\nMôi trường, sấy, bàn giao?\n\nDùng 5 Why / xương cá cùng người vận hành.", ORANGE, 16)
    d.notes("Tên lỗi được chuẩn hóa bằng phần sau dấu phân cách của mã MES, giữ nguyên outcome_detail gốc ở bảng sự kiện. Ba nhóm gồm Dent/sink mark 78, Uneven surface 56, Pinholes 55, tổng 189 trên 210 lần lỗi. Khác tiền tố 5R/7Q/8Q/9Q không bị tách thành lỗi khác khi tên giống nhau. Các câu hỏi bên phải chỉ là giả thuyết nghiệp vụ để điều tra, chưa có dữ liệu công nghệ chứng minh một nguyên nhân cụ thể.")

    # 18 — Process mining.
    d.new("06 / Process Mining", "Discovery và replay giúp khoanh vùng cần xác minh",
          "Các miner được khám phá và replay trên cùng 8 trace, nên đây là đánh giá mô tả.", ["mining/discovery_quality.csv", "mining/summary.json"])
    discovery = load("mining/discovery_quality.csv")
    d.table(42, 151, [245, 170, 170, 291], ["Mô hình", "Log fitness", "Token precision", "Diễn giải"], [
        [r.model.title()+" Miner", vn(r.log_fitness,3), vn(r.token_precision,3),
         "Phản ánh mẫu tốt hơn" if r.model == "inductive" else "Khó với vòng lặp / nhiễu"] for r in discovery.itertuples()
    ], row_h=51, size=15)
    d.card(42, 327, 429, 99, "Kiểm tra cấu trúc BPMN", "26 bước chuyển cần thiếu/bước khác để giải thích; giữ ID sự kiện để xác minh.", BLUE, 15)
    d.card(489, 327, 429, 99, "Không gán mọi sai lệch là lỗi", "Trace thiếu đầu/cuối, tên task trùng và đóng gói ngoài phạm vi đều ảnh hưởng replay.", ORANGE, 15)
    d.band("structurally_possible ≠ cả trace tuân thủ; timer 72 giờ được kiểm tra riêng bằng timestamp.")
    d.notes("Inductive Miner noise threshold 0,2 cho log fitness khoảng 0,990 và token precision 0,542; Alpha khoảng 0,501 và precision 0 trên mẫu này. Fitness cao không chứng minh tính đúng ngoài mẫu và không bằng full conformance. Kiểm tra cấu trúc duy trì tập node ứng viên cho nhãn trùng, đi qua gateway nhưng không bỏ task bắt buộc, bỏ qua guards dữ liệu. Sau deviation thì kiểm tra cục bộ tiếp, nên không được dùng tỷ lệ cặp hợp cấu trúc như fitness toàn trace. Token replay mô hình gốc được xuất riêng trong code.")

    # 19 — Forecast data design.
    d.new("07 / Thiết kế ML", "Dự báo cuối tháng t cho WIP tháng t+1",
          "Định nghĩa nhãn calendar_v1; không dùng dữ liệu tương lai làm feature.", ["forecast/feature_contract.json"])
    d.card(42, 143, 429, 139, "Nhãn 1 — Top WIP", "Top ceil(20% × số trạm) theo WIP tháng sau. Đồng hạng xử lý theo tên trạm.", BLUE, 17)
    d.card(489, 143, 429, 139, "Nhãn 2 — Ùn WIP", "WIP(t+1) > max(20; 1,5 × median WIP quá khứ). Cửa sổ t−6…t−1, ≥3 quan sát.", TEAL, 17)
    d.rect(42, 322, 505, 63, BLUE, radius=5)
    d.rect(564, 322, 354, 63, TEAL, radius=5)
    d.text(58, 334, 470, 43, "TRAIN  •  594 dòng / 11 kỳ hợp lệ\nFeature 04/2019–06/2020", 16, bold=True, color=WHITE)
    d.text(580, 334, 322, 43, "TEST  •  486 dòng / 9 kỳ\nFeature 07/2020–03/2021", 16, bold=True, color=WHITE)
    d.band("Target test: 08/2020–04/2021. Nhãn ùn WIP mới khác D5 cũ ở 98 dòng trùng phạm vi.")
    d.notes("Reindex toàn bộ lịch tháng; 08/2019 và 02/2020 giữ NaN, không nén lịch thành kỳ liền nhau. Lag1 đúng một tháng lịch, median dùng các tháng trước t và cần ít nhất 3 quan sát. Không tạo target qua một tháng thiếu. Nhãn surge mô tả mức cao so với lịch sử, không bắt buộc tăng so với ngay tháng t. Do nhãn và lag được định nghĩa lại, không so sánh trực tiếp điểm số mới với bảng Excel E1. Nhãn train muộn nhất phải quan sát được trước hoặc tại cutoff dự báo đầu tiên.")

    # 20 — Graph representation.
    d.new("07 / Graph ML", "Graph bổ sung bối cảnh tuyến, không mô phỏng thực thi BPMN",
          "Topology từ routing BOM; trọng số tự tính từ WIP lịch sử đến cutoff.", ["forecast/graph_contract.json"])
    for i, (title, body) in enumerate([
        ("Routing snapshot", "54 trạm; giữ cạnh có hỗ trợ BOM. Chưa có lịch sử phiên bản."),
        ("Walk có p/q", "Node2Vec second-order: p=1; q=0,5; 30 walk/nút, dài 20."),
        ("Vector 16 chiều", "Skip-gram; căn chỉnh Procrustes tuần tự theo thời gian."),
        ("RF + ablation", "So với WIP-only để biết graph có tạo thêm giá trị không."),
    ]):
        x = 42+i*223.5
        d.card(x, 159, 205, 228, title, body, [TEAL, BLUE, ORANGE, TEAL][i], 16)
        if i < 3:
            d.line(x+207, 271, x+220, 271, MUTED, 1.4, arrow=True)
    d.band("Kết quả graph có điều kiện: snapshot routing phải đã được biết tại thời điểm dự báo.")
    d.notes("Graph metrics giữ hướng cạnh; Node2Vec sử dụng phép chiếu vô hướng và cộng trọng số cạnh đối ứng. Trọng số là 1 + 0,5(log1p(mean WIP nguồn)+log1p(mean WIP đích)), dùng 3 tháng lịch đến t; đây không phải số lần chuyển thực. Không sử dụng D3 không rõ công thức hoặc số đếm MES tương lai. Embedding được căn giữa, chuẩn hóa và căn Procrustes với kỳ trước, không căn vào kỳ test tương lai. Một vector không mang ngữ nghĩa XOR/AND hay chứng minh trạm có công suất thiếu.")

    # 21 — Protocol.
    d.new("07 / Đánh giá ML", "Thử mô hình phức tạp phải đi cùng baseline đủ mạnh",
          "Hyperparameter cố định; đánh giá cùng split, cùng nhãn và cùng ngân sách quan sát.", ["forecast/evaluation_protocol.json"])
    d.table(42, 148, [235, 346, 295], ["Nhóm so sánh", "Cấu hình", "Câu hỏi kiểm tra"], [
        ["Quy tắc", "Majority; giữ nguyên hiện trạng", "WIP có tính ổn định mạnh không?"],
        ["RF đối chứng", "Chỉ mã trạm; WIP và lag", "Học thuộc trạm hay học tín hiệu?"],
        ["Graph ablation", "+ graph metrics; + Node2Vec; + cả hai", "Graph có giúp hơn WIP-only?"],
        ["Độ vững", "5 seed; bootstrap theo tháng", "Kết quả nhạy với random / kỳ test?"],
        ["Độ đo", "P / R / F1, ROC-AUC, AP, confusion", "Cảnh báo đúng, bỏ sót và báo nhầm?"],
    ], row_h=47, size=13.5)
    d.band("Imputer fit trên train; rework/scrap hiện tại bị loại vì chưa rõ thời điểm báo cáo có sẵn.")
    d.notes("RF có 300 cây, max_depth=6, class_weight balanced và threshold=0,5; các seed là 17,42,73,101,137. Validation nằm trong train từ 01/2020, không chọn hyperparameter để tối đa test. Code có thêm expanding-window baseline xuất riêng, không trộn với holdout cố định. Bootstrap lấy cả tháng để giữ phụ thuộc giữa trạm trong tháng; khoảng tin cậy chỉ thăm dò do ít kỳ và phụ thuộc thời gian. Độ lệch chuẩn giữa seed là biến thiên thuật toán, không phải bất định khái quát hóa.")

    # 22 — Main comparison, readable paired charts.
    d.new("08 / Kết quả ML", "Graph ML chưa vượt RF WIP trong thực nghiệm này",
          "F1 trung bình 5 seed; các quy tắc chỉ có một kết quả cố định.", ["forecast/model_comparison.csv"])
    codes = ["persistence", "RF_WIP", "RF_WIP_graph", "RF_WIP_n2v", "RF_WIP_graph_n2v"]
    names = ["Giữ nguyên", "RF WIP", "+ Graph metrics", "+ Node2Vec", "+ Graph + N2V"]
    for x, target, title in [(42, "y_surge_next", "Ùn WIP — 174 positive / 486"),
                              (489, "y_top20_next", "Top WIP — 99 positive / 486")]:
        d.text(x, 149, 429, 31, title, 18, bold=True, color=NAVY)
        values = [model[(model.target == target) & (model.model == code)].f1_mean.iloc[0] for code in codes]
        d.bars(x, 199, 429, names, values, max_value=1, row_h=42, label_w=135,
               colors=[MUTED, TEAL, BLUE, ORANGE, BLUE], value_format=lambda v: vn(v, 3))
    d.band("Top WIP có tính duy trì mạnh: persistence F1 = 0,889. Không cần ép Graph ML phải thắng.")
    d.notes("Đối với y_surge_next: persistence F1≈0,593; RF WIP≈0,608; RF WIP+graph≈0,605; +Node2Vec≈0,589; +cả hai≈0,580. Với y_top20_next: persistence≈0,889; RF WIP≈0,833; graph≈0,826; Node2Vec≈0,825; cả hai≈0,826. Đây là trung bình độ đo từng seed, không phải độ đo của xác suất ensemble. Nhãn top20 được giữ riêng với surge, không gọi cả hai là nút thắt công suất. Kết quả âm của graph là phát hiện có giá trị khi đánh giá đúng baseline.")

    # 23 — Ensemble matrix and uncertainty.
    d.new("08 / Kết quả ML", "Đọc cả báo nhầm, bỏ sót và bất định",
          "Ví dụ RF WIP ensemble 5 seed, nhãn ùn WIP; tách khỏi F1 trung bình từng seed ở slide trước.", ["forecast/ensemble_metrics_month_bootstrap.csv"])
    r = ensemble[(ensemble.target == "y_surge_next") & (ensemble.model == "RF_WIP")].iloc[0]
    d.text(83, 153, 385, 23, "DỰ BÁO", 13, bold=True, align="center", color=MUTED)
    d.text(142, 188, 142, 23, "Không ùn", 14, bold=True, align="center")
    d.text(306, 188, 142, 23, "Có ùn", 14, bold=True, align="center")
    for x, y, label, value, color in [(142, 221, "Đúng âm", r.TN, BLUE), (306, 221, "Báo nhầm", r.FP, ORANGE),
                                     (142, 323, "Bỏ sót", r.FN, RED), (306, 323, "Đúng dương", r.TP, TEAL)]:
        d.rect(x, y, 143, 88, color, radius=6)
        d.text(x+8, y+8, 127, 39, str(int(value)), 29, bold=True, color=WHITE, align="center")
        d.text(x+8, y+55, 127, 24, label, 13, color=WHITE, align="center")
    d.text(42, 243, 90, 48, "Thực tế\nkhông ùn", 13)
    d.text(42, 344, 90, 48, "Thực tế\ncó ùn", 13)
    d.card(510, 153, 408, 110, f"F1 = {vn(r.f1,3)}", f"Precision {vn(r.precision,3)} • Recall {vn(r.recall,3)}\nNgưỡng xác suất cố định: 0,5", TEAL, 17)
    d.card(510, 286, 408, 135, "Bootstrap theo 9 tháng", f"Khoảng 95% cho F1: {vn(r.f1_month_bootstrap_low,3)}–{vn(r.f1_month_bootstrap_high,3)}.\n\nÍt kỳ, có phụ thuộc thời gian: chỉ dùng như đánh giá thăm dò.", ORANGE, 16)
    d.notes("Ensemble lấy trung bình xác suất của 5 seed rồi áp ngưỡng 0,5. RF WIP có TN=236, FP=76, FN=65, TP=109; F1≈0,607. Khoảng bootstrap tháng khoảng [0,550;0,654], khác với std seed khoảng 0,008 của score riêng từng mô hình. Không tuyên bố khác biệt rất nhỏ giữa mô hình là có ý nghĩa thống kê hoặc có lợi ích vận hành. Hình cho thấy đánh đổi thật giữa bỏ sót và báo nhầm mà policy cần xem xét.")

    # 24 — Actionable alert workflow.
    d.new("08 / Ứng dụng cảnh báo", "Cảnh báo dẫn tới một lượt review có trách nhiệm",
          "Chính sách minh họa cố định dùng RF WIP; chưa chứng minh tác động đến vận hành.", ["forecast/alert_review_policy.csv"])
    for i, (head, body) in enumerate([
        ("Snapshot tháng t", "WIP + lịch sử hợp lệ"), ("RF WIP", "Xác suất ≥ 0,5"),
        ("Giới hạn alert", "Tối đa top 20% trạm"), ("Review nghiệp vụ", "Kế hoạch + SX + QC"),
    ]):
        x = 42+i*223.5
        d.card(x, 154, 205, 163, head, body, [BLUE, TEAL, ORANGE, TEAL][i], 17)
        if i < 3:
            d.line(x+207, 240, x+220, 240, MUTED, 1.4, arrow=True)
    d.card(42, 345, 429, 86, "Thông tin cần kiểm tra", "Tuổi tồn, tuyến kế tiếp, lỗi lặp lại, nguồn lực và ràng buộc QC.", BLUE, 14)
    d.card(489, 345, 429, 86, "Hành động có điều kiện", "Điều phối hoặc xử lý nguyên nhân sau xác minh; ghi lại quyết định và kết quả.", TEAL, 14)
    d.band("Dữ liệu tương lai được giữ ở bảng đánh giá; không xuất vào bảng cảnh báo minh họa.")
    d.notes("Review_alert yêu cầu xác suất ít nhất 0,5 và nằm trong ngân sách tối đa ceil(20% số trạm) của tháng; không ép đủ số cảnh báo khi xác suất thấp. Policy chọn trước RF WIP vì không phụ thuộc giả định topology lịch sử, không chọn mô hình dựa trên điểm test tốt nhất. File alert_review_policy bỏ actual và outcome tương lai. Trước khi áp dụng cần quyết định ai có quyền điều phối và giới hạn thay đổi, rồi đo precision, khối lượng review và kết quả sau can thiệp.")

    # 25 — To-Be mechanisms.
    d.new("09 / Thiết kế To-Be", "Ba cơ chế cải tiến, ba chỉ số đánh giá riêng",
          "Giữ cổng QC và thời gian ổn định bắt buộc; ghi rõ chi phí và giả định.", ["scenarios/scenario_contract.json"])
    for yy, label, tasks in [(154, "AS-IS", ["Chuẩn bị bề mặt", "Sơn", "QC6"]),
                              (244, "TO-BE", ["Chuẩn bị bề mặt", "Checklist 2 phút", "Sơn", "QC6"])]:
        d.text(42, yy+22, 90, 30, label, 16, bold=True, color=TEAL if label == "TO-BE" else MUTED)
        for i, task in enumerate(tasks):
            xx = 151+i*187
            active = "Checklist" in task
            d.rect(xx, yy, 167, 64, TEAL if active else WHITE, stroke=None, radius=6)
            d.text(xx+10, yy+19, 147, 37, task, 17, bold=True, color=WHITE if active else INK, align="center")
            if i < len(tasks)-1:
                d.line(xx+169, yy+32, xx+184, yy+32, MUTED, 1.4, arrow=True)
    for i, (title, body, color) in enumerate([
        ("Dispatch • Kế hoạch", "Thử điều phối vào chà nhám; giảm gap là tham số kịch bản.", BLUE),
        ("Checklist • SX + QC", "Giữ QC6; đo lại lỗi và workload sau thay đổi điều kiện sơn.", TEAL),
        ("Ổn định • MES + QC", "Giữ đủ 72h; bố trí kiểm tra sớm sau mở khóa.", ORANGE),
    ]):
        d.card(42+i*298, 341, 280, 86, title, body, color, 13)
    d.band("Checklist giả định tốn 2 phút/lượt sơn; hiệu quả giảm lỗi cần được thử nghiệm thực tế.")
    d.notes("To-Be không chỉ thêm một hình cảnh báo ML. Mỗi cơ chế có chủ sở hữu và phép đo phù hợp. Code tạo hai BPMN projection riêng cho chà nhám–sơn theo lớp, có nhánh sửa/đạt/hủy và chuyển lớp; To-Be thêm checklist. Policy điều phối chưa có công suất để hiệu chỉnh nên mức giảm gap được khảo sát như giả định. Cơ chế timer QC7 độc lập và không được rút ngắn 72h. Chưa cộng các estimand thành lợi ích toàn nhà máy.")

    # 26 — Gap replay.
    d.new("09 / Đánh giá To-Be", "Replay gap: lợi ích phụ thuộc giả định và case dài",
          "Giữ nguyên timeline serial quan sát, nén gap vào chà nhám và tính thêm chi phí checklist.", ["scenarios/gap_replay_summary.csv"])
    g = gap[gap.checklist_minutes_per_observed_paint.eq(2)].sort_values("assumed_gap_reduction")
    d.table(42, 150, [230, 207, 222, 217], ["Giảm gap giả định", "TB tiết kiệm / case", "Trung vị / case", "Khoảng TB sau replay"], [
        [vn(r.assumed_gap_reduction*100)+"%", vn(r.mean_saving_hours,2)+" giờ", vn(r.median_saving_hours,2)+" giờ",
         vn(r.mean_scenario_hours,1)+" giờ"] for r in g.itertuples()
    ], row_h=49, size=14, highlights=[2])
    d.band("Kịch bản 25%: TB ≈ 189,86 giờ nhưng trung vị ≈ 33,49 giờ/case → ảnh hưởng mạnh của case dài.")
    d.notes("Baseline khoảng quan sát trung bình là 1.908,12 giờ/case. Checklist 2 phút mỗi lượt sơn làm tăng trung bình 0,9375 giờ/case khi chưa giảm gap. Giả định giảm gap 25% cho saving TB 189,86 giờ và median 33,49 giờ. Khoảng quan sát sau replay vẫn không phải Cycle Time hoàn thành đã đo, vì trace có thể bị cắt cuối. Đây là phép tính timeline tuần tự có điều kiện, không mô phỏng hàng đợi tài nguyên và không chứng minh nguồn lực đủ để đạt mức giảm đó.")

    # 27 — Workload Monte Carlo.
    d.new("09 / Đánh giá To-Be", "Giảm làm lại được lượng hóa thành workload định mức",
          "10.000 lượt mô phỏng; mỗi lượt gồm 100 SP, ghép As-Is/To-Be bằng cùng số ngẫu nhiên.", ["scenarios/repeat_workload_sensitivity.csv"])
    selected = workload[workload.assumed_relative_reduction_in_mean_repeats.eq(.25)]
    for i, r in enumerate(selected.itertuples()):
        x = 42+i*447
        d.rect(x, 153, 429, 232, WHITE, radius=8)
        d.text(x+18, 170, 390, 32, r.stage.upper(), 18, color=TEAL if i == 0 else BLUE, bold=True)
        d.text(x+18, 216, 390, 54, vn(r.analytic_expected_saving_minutes/60,2)+" giờ / 100 SP", 30,
               bold=True, color=TEAL if i == 0 else BLUE)
        d.text(x+18, 291, 390, 73, f"Giả định giảm R: 25%\nR tổng thể: {vn(r.population_mean_repeats_per_completed_product,4)}\nĐịnh mức proxy: {vn(r.standard_minutes_per_repeat_proxy,1)} phút/lượt", 17)
    d.band("Đây là giờ định mức làm lại tiết kiệm; không phải giờ rút ngắn Cycle Time hoặc tăng throughput.")
    d.notes("Geometric-equivalent dùng r=R/(1+R), phù hợp kỳ vọng số lần làm lại được cung cấp nhưng chưa chứng minh thực tế các lần lặp độc lập. Kỳ vọng giải tích tiết kiệm =100×R×T×mức giảm R. Với 25%: chà nhám 898,95 phút≈14,98 giờ, sơn 99,627 phút≈1,66 giờ. T lấy median định mức operation từ mẫu MES, khác nguồn với R tổng thể nên chỉ là proxy. Không tính lớp kế tiếp vào repeated workload. Quantile mô phỏng là biến thiên lấy mẫu có điều kiện, không phải CI nhân quả hay bất định tham số.")

    # 28 — QC timer.
    d.new("09 / Đánh giá To-Be", "QC7: đủ 72 giờ rồi giảm chậm bố trí kiểm tra",
          "Chỉ quan sát hai cặp chờ tại CASE05 / SP04; chưa khái quát cho mọi tuyến.", ["scenarios/qc7_timer_dispatch_replay.csv"])
    for i, r in enumerate(timers.itertuples()):
        y = 154+i*129
        d.text(42, y, 155, 29, f"Lần chờ {i+1}", 18, bold=True)
        d.text(207, y, 690, 24,
               f"Gốc: {vn(r.timer_hours,2)}h ổn định + {vn(r.post_unlock_gap_hours,2)}h đến QC7", 16)
        span = 510
        d.rect(207, y+39, span*r.baseline_block_to_qc_hours/145, 22, ORANGE, radius=3)
        d.text(745, y+36, 173, 28, vn(r.baseline_block_to_qc_hours,2)+"h gốc", 16, color=ORANGE, bold=True)
        d.rect(207, y+76, span*r.scenario_block_to_qc_hours/145, 22, TEAL, radius=3)
        d.text(745, y+73, 173, 28, vn(r.scenario_block_to_qc_hours,0)+"h giả định", 16, color=TEAL, bold=True)
    d.band("To-Be: 72h bắt buộc + dispatch giả định 2h → tiết kiệm có điều kiện 18,82h và 66,08h.")
    d.notes("Dữ liệu cho thấy cặp BLOCK/UNBLOCK 71,037h và 71,844h, đều ngắn hơn 72h; sau mở khóa đến Start QC7 thêm khoảng 21,786h và 68,240h. Replay nâng timer lên ít nhất 72h và đặt dispatch giả định tối đa 2h, nên cả hai còn 74h. Mức saving tương ứng 18,823h và 66,084h, tính từ BLOCK đến Start QC7, không phải đến khi hoàn tất QC. Khả năng đạt dispatch 2h chưa được hiệu chỉnh bằng công suất/ca làm, cần thử nghiệm.")

    # 29 — Reproducibility, limitations, pilot.
    d.new("10 / Kiểm chứng & giới hạn", "Có thể tái lập kết quả; tác động thực tế còn cần pilot",
          "Tái lập số liệu là điều kiện cần, chưa thay thế kiểm chứng ngoài hiện trường.", ["metadata/verification.json", "metadata/manifest.json"])
    d.metric(42, 146, 280, "17", "Test đã pass", "Thời gian, leakage, graph, mô phỏng")
    d.metric(340, 146, 280, str(verification["saved_models_reproduced"]), "Model được nạp kiểm tra", "Tái tạo xác suất dự báo đã lưu", BLUE)
    d.metric(638, 146, 280, vn(verification["prediction_rows_checked"]), "Dòng dự đoán kiểm tra", "64 nhóm độ đo được tính lại", ORANGE)
    d.card(42, 303, 429, 130, "Giới hạn phải giữ khi trình bày", "Mẫu MES thiên lệch; chưa có ready_time/công suất; topology chưa có version; 9 kỳ test; kịch bản chưa có kiểm chứng nhân quả.", ORANGE, 15)
    d.card(489, 303, 429, 130, "Bước tiếp theo ở doanh nghiệp", "Bổ sung ready_time, ca và nguồn lực; xác minh lỗi; pilot nhỏ có nhóm so sánh phù hợp; đo thời gian, chất lượng và workload.", TEAL, 15)
    d.notes("uv.lock, config và manifest lưu phiên bản, hash input/code; script verify kiểm tra lại mọi nhóm metrics, 12 model seed42 và các bất biến mô phỏng. Các test còn có regression số học Word2Vec cho dependency Gensim được khóa commit sửa lỗi upstream. Thực nghiệm chỉ kiểm tra tháng tương lai tại trạm đã biết, chưa kiểm tra nhà máy mới hoặc trạm mới. Kế hoạch pilot cần mục tiêu và tiêu chí chất lượng do doanh nghiệp xác nhận; không đặt mức giảm giả định làm kết quả cam kết.")

    # 30 — Final synthesis with references in speaker notes.
    d.new("Kết luận", "Từ sơ đồ BPMN đến cải tiến có thể kiểm chứng", dark=True)
    d.text(48, 42, 850, 24, "KẾT LUẬN & TRAO ĐỔI", 12, bold=True, color="A1C9D0")
    d.text(48, 100, 850, 88, "Ưu tiên chà nhám–sơn.\nGiữ baseline. Kiểm chứng To-Be.", 34, bold=True, color=WHITE)
    for y, number, title, body in [
        (224, "01", "Bằng chứng", "WIP + lượt lặp + lỗi theo lớp cùng chỉ về vùng cần điều tra."),
        (292, "02", "Dự báo", "RF WIP hữu ích ở mức thăm dò; graph chưa tạo lợi thế trong mẫu."),
        (360, "03", "Hành động", "Checklist, dispatch và QC7 cần pilot với dữ liệu nguồn lực thực tế."),
    ]:
        d.text(49, y, 57, 42, number, 27, bold=True, color="64CBC3")
        d.text(119, y, 730, 25, title, 19, bold=True, color=WHITE)
        d.text(119, y+29, 762, 28, body, 16, color="CBDCE6")
    d.text(49, 450, 840, 25, "Thành viên: [Bổ sung sau]   •   Câu hỏi & thảo luận", 15, color=WHITE)
    d.notes("Tổng hợp: nhóm đã xây chuỗi phân tích từ dữ liệu đã ẩn danh đến các bằng chứng điểm nghẽn, so sánh dự báo và kịch bản cải tiến có giả định. Không đồng nhất top WIP với bottleneck công suất, không biến kết quả mô phỏng thành thành tích công ty.\n\nTài liệu nền: Báo cáo tổng hợp 6 buổi của Võ Minh Trí; mô tả quy trình và BPMN Hình 1–5; CongTyX_Event log.xlsx. Tham chiếu phương pháp: OMG BPMN 2.0.2; Dumas et al., Fundamentals of Business Process Management (2018); van der Aalst, Process Mining (2016); PM4Py (https://github.com/process-intelligence-solutions/pm4py); Grover & Leskovec, node2vec (2016), https://arxiv.org/abs/1607.00653. Toàn bộ số liệu slide được lấy từ src/outputs; xem slide_sources.json để truy nguyên và hash. Người trình bày có thể mở demo_executed.ipynb để minh họa cách đi từ bảng kết quả về event_id/dòng A1.")

    d.save()
    render_previews()
    validate_deck()


def render_previews():
    directory = HERE / "preview"
    directory.mkdir(exist_ok=True)
    pdf = fitz.open(HERE / "BPMN_CongTyX_30slides.pdf")
    thumbnails = []
    for index, page in enumerate(pdf):
        pix = page.get_pixmap(matrix=fitz.Matrix(1.25, 1.25), alpha=False)
        path = directory / f"slide_{index+1:02d}.png"
        pix.save(path)
        image = Image.open(path).convert("RGB")
        image.thumbnail((384, 216))
        thumb = Image.new("RGB", (400, 240), "#E1E7EE")
        thumb.paste(image, (8, 4))
        ImageDraw.Draw(thumb).text((10, 221), f"{index+1:02d}", fill="#20384B")
        thumbnails.append(thumb)
    for sheet, start in enumerate(range(0, len(thumbnails), 10), 1):
        contact = Image.new("RGB", (2000, 480), "#E1E7EE")
        for i, thumbnail in enumerate(thumbnails[start:start+10]):
            contact.paste(thumbnail, ((i % 5)*400, (i//5)*240))
        contact.save(directory / f"contact_{sheet}.png")


def validate_deck():
    pptx = Presentation(HERE / "BPMN_CongTyX_30slides.pptx")
    pdf = fitz.open(HERE / "BPMN_CongTyX_30slides.pdf")
    assert len(pptx.slides) == len(pdf) == 30
    outside = []
    for index, slide in enumerate(pptx.slides, 1):
        assert slide.has_notes_slide
        assert len(slide.notes_slide.notes_text_frame.text) > 100
        for shape in slide.shapes:
            if shape.left < -2 or shape.top < -2 or shape.left+shape.width > pptx.slide_width+2 or shape.top+shape.height > pptx.slide_height+2:
                outside.append({"slide": index, "shape": shape.name})
    if outside:
        raise ValueError(f"Shapes outside slide: {outside}")
    # Make sure Vietnamese text was embedded/extracted from the PDF, not replaced by missing glyphs.
    text = "\n".join(page.get_text() for page in pdf)
    assert "chà nhám" in text.lower() and "kiểm chứng" in text.lower()
    assert "\ufffd" not in text
    result = {"slides": 30, "aspect_ratio": "16:9", "speaker_notes": 30,
              "out_of_bounds_shapes": outside, "text_fit": "measured before writing; no overflow accepted",
              "preview_pages": len(pdf), "vietnamese_text_extraction": "passed",
              "pptx_sha256": sha(HERE / "BPMN_CongTyX_30slides.pptx"),
              "pdf_sha256": sha(HERE / "BPMN_CongTyX_30slides.pdf")}
    (HERE / "validation.json").write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    build()
